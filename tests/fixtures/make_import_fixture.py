"""
Génère tests/fixtures/import_complexe.xlsx : classeur de test de l'import Excel → Word.

    python3 tests/fixtures/make_import_fixture.py

Le fichier produit par openpyxl ne contient pas les résultats des formules ; il est ensuite recalculé
et réenregistré par LibreOffice (import_complexe.xlsx) :

    soffice --headless --convert-to xlsx:"Calc MS Excel 2007 XML" --outdir tests/fixtures tests/fixtures/import_complexe_sans_calcul.xlsx
    mv tests/fixtures/import_complexe_sans_calcul.xlsx ... (voir le script : il fait tout)
"""
import datetime
import os
import shutil
import subprocess
import tempfile

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))

thin = Side(style="thin", color="000000")
medium = Side(style="medium", color="1F3864")
double = Side(style="double", color="000000")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
header_fill = PatternFill("solid", fgColor="4472C4")
light_fill = PatternFill("solid", fgColor="D9E1F2")
white_bold = Font(name="Calibri", size=11, bold=True, color="FFFFFF")


def ventes(ws):
    ws["A1"] = "Tableau 1 : Ventes trimestrielles par région"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True)
    headers = ["Région", "T1", "T2", "T3", "T4", "Total", "Évolution", "Code interne"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = white_bold
        c.fill = header_fill
        c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data = [
        ("Île-de-France", 125430.5, 131200.25, 128900, 142750.75, 0.0825, "IDF-01"),
        ("Auvergne-Rhône-Alpes", 98765.4, 97500, 101230.1, 99870.6, -0.0112, "ARA-02"),
        ("Nouvelle-Aquitaine", 67890, 70120.35, 69980.5, 73450.25, 0.0461, "NAQ-03"),
        ("Occitanie", 54320.75, 52100, 51980.4, 50210.9, -0.0759, "OCC-04"),
        ("Hauts-de-France", 61234.5, 63210.8, 64890, 66120.45, 0.0312, "HDF-05"),
    ]
    currency = '#,##0.00\\ "€"'
    for r, row in enumerate(data, start=4):
        ws.cell(row=r, column=1, value=row[0]).border = box
        for j in range(4):
            c = ws.cell(row=r, column=2 + j, value=row[1 + j])
            c.number_format = currency
            c.border = box
        t = ws.cell(row=r, column=6, value="=SUM(B{0}:E{0})".format(r))
        t.number_format = currency
        t.border = box
        t.font = Font(bold=True)
        e = ws.cell(row=r, column=7, value=row[5])
        e.number_format = '0.0%;[Red]-0.0%'
        e.border = box
        ws.cell(row=r, column=8, value=row[6])
        if r % 2 == 1:
            for j in range(1, 8):
                ws.cell(row=r, column=j).fill = light_fill
    # Ligne masquée (ne doit pas être importée).
    ws.cell(row=9, column=1, value="Ligne masquée").border = box
    ws.row_dimensions[9].hidden = True
    total_row = 10
    ws.cell(row=total_row, column=1, value="Total").font = Font(bold=True)
    for j in range(2, 7):
        col = "BCDEF"[j - 2]
        c = ws.cell(row=total_row, column=j, value="=SUM({0}4:{0}8)".format(col))
        c.number_format = currency
        c.font = Font(bold=True)
    for j in range(1, 8):
        ws.cell(row=total_row, column=j).border = Border(left=thin, right=thin, top=double, bottom=medium)
    ws.cell(row=12, column=1, value="Source : service commercial, données provisoires au 31/12/2024 (hors taxes).")
    ws.cell(row=12, column=1).font = Font(italic=True, size=9, color="595959")
    ws.column_dimensions["A"].width = 24
    for col in "BCDE":
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 11
    ws.column_dimensions["H"].hidden = True
    ws.row_dimensions[3].height = 30
    ws.conditional_formatting.add("G4:G8", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", bgColor="FFC7CE"), font=Font(color="9C0006")))


def planning(ws):
    headers = ["Phase", "Tâche", "Début", "Fin", "Durée", "Terminé", "Avancement", "Remarque"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFE699")
        c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center", text_rotation=90 if h == "Terminé" else 0)
    rows = [
        ("Phase 1", "Cadrage du projet", datetime.datetime(2024, 1, 8, 9, 0), datetime.datetime(2024, 1, 19, 17, 30), 1.5, True, 1, "Validé en comité"),
        (None, "Recueil des besoins", datetime.datetime(2024, 1, 22, 9, 0), datetime.datetime(2024, 2, 16, 18, 0), 2.25, True, 1, "Ateliers avec les utilisateurs\net les métiers (4 sessions)"),
        (None, "Rédaction des spécifications", datetime.datetime(2024, 2, 19), datetime.datetime(2024, 3, 29), 10.75, False, 0.85, None),
        ("Phase 2", "Développement", datetime.datetime(2024, 4, 1), datetime.datetime(2024, 7, 12), 30.5, False, 0.4, "=1/0"),
        (None, "Recette", datetime.datetime(2024, 9, 2), datetime.datetime(2024, 10, 11), 12, False, 0, None),
    ]
    for r, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.border = box
        ws.cell(row=r, column=3).number_format = "dddd d mmmm yyyy"
        ws.cell(row=r, column=4).number_format = "dd/mm/yyyy hh:mm"
        ws.cell(row=r, column=5).number_format = "[h]:mm"
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=7).number_format = "0%"
        ws.cell(row=r, column=8).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=2).alignment = Alignment(indent=1, vertical="center")
    ws.merge_cells("A2:A4")
    ws.merge_cells("A5:A6")
    for a in ("A2", "A5"):
        ws[a].alignment = Alignment(horizontal="center", vertical="center")
        ws[a].font = Font(bold=True, color="C00000")
    rich = CellRichText([TextBlock(InlineFont(b=True, color="C00000"), "Attention : "), "livrable en retard de ", TextBlock(InlineFont(i=True, u="single"), "deux semaines")])
    ws.cell(row=7, column=1, value="Note").font = Font(bold=True)
    ws.cell(row=7, column=2, value=rich)
    ws.merge_cells("B7:H7")
    ws.cell(row=8, column=1, value="Montant")
    ws.cell(row=8, column=2, value=-1234.5).number_format = '#,##0.00;[Red](#,##0.00)'
    ws.cell(row=8, column=3, value=0.000123).number_format = "0.00E+00"
    ws.cell(row=8, column=4, value=3.14159265).number_format = "# ??/??"
    ws.cell(row=8, column=5, value=1234567890123).number_format = "General"
    ws.cell(row=8, column=6, value=False)
    ws.cell(row=8, column=7, value="0612345678").number_format = "@"
    ws.cell(row=8, column=8, value=0.1 + 0.2)
    widths = {"A": 11, "B": 30, "C": 26, "D": 17, "E": 9, "F": 5, "G": 12, "H": 32}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.row_dimensions[1].height = 62


def tableau_excel(ws):
    ws.append(["Produit", "Catégorie", "Prix unitaire", "Quantité", "Montant"])
    items = [("Stylo", "Fournitures", 1.2, 150), ("Cahier", "Fournitures", 2.5, 80), ("Agrafeuse", "Matériel", 12.9, 12),
             ("Écran 24 pouces", "Informatique", 189.99, 5), ("Clavier", "Informatique", 24.5, 10)]
    for i, (p, cat, prix, q) in enumerate(items, start=2):
        ws.append([p, cat, prix, q, "=C{0}*D{0}".format(i)])
        ws.cell(row=i, column=3).number_format = '#,##0.00\\ "€"'
        ws.cell(row=i, column=5).number_format = '#,##0.00\\ "€"'
    table = Table(displayName="Achats", ref="A1:E6")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)
    for col, w in zip("ABCDE", (20, 15, 14, 10, 14)):
        ws.column_dimensions[col].width = w


def main():
    wb = Workbook()
    ventes(wb.active)
    wb.active.title = "Ventes 2024"
    planning(wb.create_sheet("Planning"))
    tableau_excel(wb.create_sheet("Achats"))
    hidden = wb.create_sheet("Paramètres")
    hidden["A1"] = "Feuille masquée"
    hidden.sheet_state = "hidden"
    wb.create_sheet("Vide")
    large = wb.create_sheet("Très large")
    for c in range(1, 71):
        large.cell(row=1, column=c, value="C{0}".format(c))

    raw = os.path.join(HERE, "import_complexe_sans_calcul.xlsx")
    wb.save(raw)

    # Recalcul par LibreOffice pour enregistrer les résultats des formules.
    out = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, raw],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shutil.move(os.path.join(out, "import_complexe_sans_calcul.xlsx"), os.path.join(HERE, "import_complexe.xlsx"))
    shutil.rmtree(out, ignore_errors=True)


if __name__ == "__main__":
    main()
