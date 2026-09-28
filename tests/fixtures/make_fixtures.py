"""
Génère les fichiers XML de test utilisés par WordXmlTableParserTests.

- medium_shading_merged.flat.xml : paquet « Flat OPC », comme Range.WordOpenXML (Word 2010+),
  contenant document.xml et styles.xml d'un document créé avec python-docx
  (style « Medium Shading 1 Accent 1 », fusions horizontale et verticale, trame et bordure directes).
- medium_shading_merged.docx     : le même document (pour inspection manuelle).
- libreoffice_word2003.xml       : conversion LibreOffice au format « Word 2003 XML » (espace de noms
  WordprocessingML 2003, largeurs décimales, trames « solid ») : sert à vérifier la robustesse de l'analyse.

Usage : python3 make_fixtures.py   (nécessite python-docx ; LibreOffice Writer pour le fichier 2003)
"""
import copy
import os
import subprocess
import zipfile

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, edge, val, sz, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    b = OxmlElement("w:" + edge)
    b.set(qn("w:val"), val)
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "0")
    b.set(qn("w:color"), color)
    borders.append(b)


def build_document():
    d = docx.Document()
    p = d.add_paragraph("Tableau 1 : Ventes par région")
    table = d.add_table(rows=5, cols=4)
    table.style = d.styles["Medium Shading 1 Accent 1"]

    # Options de style : ligne d'en-tête, dernière ligne, première colonne, lignes à bandes.
    tbl_pr = table._tbl.tblPr
    look = tbl_pr.find(qn("w:tblLook"))
    if look is None:
        look = OxmlElement("w:tblLook")
        tbl_pr.append(look)
    for k in list(look.attrib):
        del look.attrib[k]
    look.set(qn("w:val"), "04E0")
    look.set(qn("w:firstRow"), "1")
    look.set(qn("w:lastRow"), "1")
    look.set(qn("w:firstColumn"), "1")
    look.set(qn("w:lastColumn"), "0")
    look.set(qn("w:noHBand"), "0")
    look.set(qn("w:noVBand"), "1")

    data = [
        ["Région", "Ventes", "", "Évolution"],
        ["Nord", "1 200", "1 350", "12,5 %"],
        ["", "800", "760", "-5 %"],
        ["Sud", "2 000", "2 100", "5 %"],
        ["Total", "4 000", "4 210", "5,25 %"],
    ]
    for r, row in enumerate(data):
        for c, text in enumerate(row):
            if text:
                table.cell(r, c).text = text

    table.cell(0, 1).merge(table.cell(0, 2))      # fusion horizontale
    table.cell(1, 0).merge(table.cell(2, 0))      # fusion verticale
    set_cell_shading(table.cell(3, 3), "FFFF00")  # trame directe
    set_cell_border(table.cell(1, 1), "bottom", "single", 18, "FF0000")  # bordure directe

    d.add_paragraph("Source : service commercial.")
    return d


def flat_opc(docx_path):
    with zipfile.ZipFile(docx_path) as z:
        document = etree.fromstring(z.read("word/document.xml"))
        styles = etree.fromstring(z.read("word/styles.xml"))
    pkg = "http://schemas.microsoft.com/office/2006/xmlPackage"
    root = etree.Element("{%s}package" % pkg, nsmap={"pkg": pkg})
    for name, content_type, element in [
        ("/word/document.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml", document),
        ("/word/styles.xml", "application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml", styles),
    ]:
        part = etree.SubElement(root, "{%s}part" % pkg)
        part.set("{%s}name" % pkg, name)
        part.set("{%s}contentType" % pkg, content_type)
        data = etree.SubElement(part, "{%s}xmlData" % pkg)
        data.append(copy.deepcopy(element))
    xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return xml.replace(b"?>", b'?>\n<?mso-application progid="Word.Document"?>', 1)


def main():
    docx_path = os.path.join(HERE, "medium_shading_merged.docx")
    build_document().save(docx_path)
    with open(os.path.join(HERE, "medium_shading_merged.flat.xml"), "wb") as f:
        f.write(flat_opc(docx_path))

    try:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", 'xml:MS Word 2003 XML', "--outdir", HERE, docx_path],
            check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
        os.replace(os.path.join(HERE, "medium_shading_merged.xml"), os.path.join(HERE, "libreoffice_word2003.xml"))
    except (OSError, subprocess.SubprocessError) as e:
        print("LibreOffice Writer indisponible, libreoffice_word2003.xml non régénéré :", e)


if __name__ == "__main__":
    main()
